"""Convert Review_2_Literature_Survey.md into a formally formatted Microsoft Word (.docx) document."""

import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def build_docx(md_path: str, docx_path: str):
    doc = docx.Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    for line in lines:
        raw = line.rstrip()

        # Check for Markdown table line
        if raw.startswith('|') and raw.endswith('|'):
            in_table = True
            table_lines.append(raw)
            continue
        elif in_table:
            # End of table, process it
            in_table = False
            parse_and_add_table(doc, table_lines)
            table_lines = []

        if not raw.strip():
            continue

        # Headings
        if raw.startswith('# '):
            p = doc.add_heading(raw[2:], level=1)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.name = 'Calibri Light'
                r.font.size = Pt(18)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 45, 95)
        elif raw.startswith('## '):
            p = doc.add_heading(raw[3:], level=2)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = 'Calibri Light'
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(25, 70, 130)
        elif raw.startswith('### '):
            p = doc.add_heading(raw[4:], level=3)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = RGBColor(40, 90, 150)
        elif raw.startswith('---'):
            continue
        elif raw.startswith('* ') or raw.startswith('- '):
            # Bullet point
            bullet_text = raw[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, bullet_text)
        elif re.match(r'^\d+\.\s', raw):
            # Numbered list
            m = re.match(r'^\d+\.\s(.*)', raw)
            text = m.group(1).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, text)
        elif raw.startswith('> '):
            # Blockquote
            quote_text = raw[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(quote_text)
            r.font.italic = True
            r.font.color.rgb = RGBColor(70, 70, 70)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, raw)

    if in_table and table_lines:
        parse_and_add_table(doc, table_lines)

    doc.save(docx_path)
    print(f"[*] Successfully built Word document at: {docx_path}")


def parse_and_add_table(doc, table_lines):
    # Filter out separator lines (|---|---|)
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip('|').split('|')]
        # Check if separator
        if all(re.match(r'^:?-+:?$', c) for c in cells if c):
            continue
        rows.append(cells)

    if not rows:
        return

    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows):
        for c_idx, cell_value in enumerate(row):
            if c_idx >= n_cols:
                continue
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05

            clean_val = cell_value.replace('<br>', '\n').replace('\\n', '\n')
            add_formatted_runs(p, clean_val)

            # Styling header vs data rows
            if r_idx == 0:
                set_cell_background(cell, "1F4E79")  # Dark blue
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F2F5F9" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                for run in p.runs:
                    run.font.size = Pt(9)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_formatted_runs(paragraph, text):
    # Basic markdown bold & italic parser
    # Match **bold** or *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            r = paragraph.add_run(token[2:-2])
            r.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            r = paragraph.add_run(token[1:-1])
            r.font.italic = True
        else:
            paragraph.add_run(token)


if __name__ == "__main__":
    build_docx("docs/Review_2_Literature_Survey.md", "docs/Review_2_Literature_Survey.docx")
